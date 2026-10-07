import csv, re
from datetime import date, datetime
from io import BytesIO, TextIOWrapper
from typing import Any
from openpyxl import load_workbook
import fitz
from ..schemas.employee import EmployeeCreate
from ..utils.logger import get_logger

logger = get_logger(__name__)
REQUIRED = {"first_name","last_name","email","phone","department","designation","salary","status","joining_date"}

class FileProcessor:
    def _normalize(self, row: dict[str, Any]) -> dict[str, Any]:
        return {str(k).strip().lower().replace(" ", "_"): v for k,v in row.items() if k is not None}

    def _valid(self, row: dict[str, Any]) -> EmployeeCreate:
        normalized = self._normalize(row)
        missing = REQUIRED - normalized.keys()
        if missing: raise ValueError(f"Missing columns: {sorted(missing)}")
        if isinstance(normalized["joining_date"], (datetime, date)):
            normalized["joining_date"] = normalized["joining_date"].isoformat()
        return EmployeeCreate(**{k: normalized[k] for k in REQUIRED})

    def csv_import(self, content: bytes) -> tuple[list[EmployeeCreate], int]:
        good, failed = [], 0
        reader = csv.DictReader(TextIOWrapper(BytesIO(content), encoding="utf-8-sig", newline=""))
        for line, row in enumerate(reader, start=2):
            if not any(row.values()): continue
            try: good.append(self._valid(row))
            except (ValueError, TypeError) as exc:
                failed += 1; logger.warning("Invalid CSV row %s: %s", line, exc)
        return good, failed

    def excel_import(self, content: bytes) -> tuple[list[EmployeeCreate], int]:
        good, failed = [], 0
        wb = load_workbook(BytesIO(content), data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows: return [], 0
        headers = [str(h).strip() if h is not None else "" for h in rows[0]]
        for idx, values in enumerate(rows[1:], start=2):
            if not any(v is not None and str(v).strip() for v in values): continue
            try: good.append(self._valid(dict(zip(headers, values))))
            except (ValueError, TypeError) as exc:
                failed += 1; logger.warning("Invalid Excel row %s: %s", idx, exc)
        return good, failed

    def pdf_import(self, content: bytes) -> dict[str, str]:
        text = ""
        with fitz.open(stream=content, filetype="pdf") as doc:
            text = "\n".join(page.get_text() for page in doc)
        patterns = {
            "first_name": r"First\s*Name\s*:\s*(.+)",
            "last_name": r"Last\s*Name\s*:\s*(.+)",
            "email": r"Email\s*:\s*([^\s]+)",
            "department": r"Department\s*:\s*(.+)",
        }
        result={}
        for key, pattern in patterns.items():
            match=re.search(pattern, text, re.I)
            if match: result[key]=match.group(1).strip()
            else: logger.warning("PDF field not matched: %s", key)
        return result
